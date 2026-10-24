#!/usr/bin/env python3
"""Export measured directional predictions using beginner-readable headings."""
import json,math,shutil,sys,zipfile
from collections import Counter
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import xlsxwriter
import openpyxl
from btc_lab.data import write_json,sha256,utc_now
from indicator_lab.data import ROOT
from directional_study.run import request,verify_inputs
OUT=ROOT/'reports/selected_directional';DEST=ROOT/'BTC_UP_DOWN_ACCURACY.xlsx'
FIELDS={
'id':('ID','The indicator code from your supplied list.'),
'name':('Indicator','The readable indicator name. Its forecast comes from the same frozen trained model, not automatically from the indicator alone.'),
'requested_subset_rank':('Subset rank','Rank by Jan–Sep 2026 overall accuracy among your requested indicators on this chart. Not ranked by profit or independent future evidence.'),
'timeframe_minutes':('Chart minutes','5, 15 or 60 minutes. Forecast the next close on that matching chart.'),
'period':('Period','validation = 2025, used for window selection; retrospective = Jan–Sep 2026, previously examined prices.'),
'overall_accuracy':('Overall correct','Correct UP and DOWN guesses divided by all scored guesses. Ignores trading costs.'),
'up_accuracy':('Correct UP guesses','Of all guesses saying UP, the share that were right: TP/(TP+FP). N/A if UP was never guessed.'),
'up_calls':('UP guesses','How many scored guesses said UP. A high percentage on very few guesses is weak evidence.'),
'down_accuracy':('Correct DOWN guesses','Of all guesses saying DOWN, the share that were right: TN/(TN+FN). N/A if DOWN was never guessed.'),
'down_calls':('DOWN guesses','How many scored guesses said DOWN. Never assume that all indicators forecast both directions.'),
'forecasts':('Scored guesses','The number of guesses with known, unequal next closes and usable features. Not a trade count.'),
'up_moves_found':('UP moves found','Of all actual UP moves, the share correctly called UP: TP/(TP+FN). Different denominator from Correct UP guesses.'),
'down_moves_found':('DOWN moves found','Of all actual DOWN moves, the share correctly called DOWN: TN/(TN+FP). Different denominator from Correct DOWN guesses.'),
'balanced_accuracy':('Balanced accuracy','Average of UP moves found and DOWN moves found. A model that always guesses one direction typically scores 50% here.'),
'predicted_up_share':('UP guess share','What percentage of all scored guesses said UP. Not its success rate.'),
'actual_up_share':('Actual UP share','What percentage of actual scored next closes were higher. Shows the class balance.'),
'up_ci_low':('UP interval low','Lower end of a descriptive 95% Wilson interval for Correct UP guesses. Assumes independent samples; not a future guarantee.'),
'up_ci_high':('UP interval high','Upper end of that descriptive UP precision interval. Read together with UP interval low and UP guess count.'),
'down_ci_low':('DOWN interval low','Lower end of a descriptive 95% Wilson interval for Correct DOWN guesses. Market predictions can be correlated.'),
'down_ci_high':('DOWN interval high','Upper end of that descriptive DOWN precision interval. Read with DOWN interval low and DOWN guess count.'),
'coverage':('Scoring coverage','Share of otherwise eligible targets that could be scored for this model.'),
'warnings':('Warnings','Flags for fewer than 100 guesses in a direction, rare patterns, absent events or special input interpretation. These are warnings, not proof of failure or reliability.'),
'pattern_fired_count':('Pattern occurrences','How many completed candles actually triggered the candlestick pattern in this period, including unscored tied outcomes.'),
'pattern_scored_count':('Pattern events scored','How many triggered candlestick events had scorable predictions. N/A for continuous indicators. Zero means no event accuracy can be measured.'),
'pattern_flat_targets':('Pattern tied outcomes','Pattern occurrences whose next close equaled the current close. Excluded from direction accuracy.'),
'pattern_other_unscored':('Other unscored events','Triggered events not scored for another eligibility reason. Not counted as wrong or correct.'),
'pattern_few_scored':('Few pattern events','TRUE if fewer than 100 pattern occurrences could be scored. No guarantee of reliability above 100.'),
'rare_training_events':('Rare training pattern','TRUE if the pattern had fewer than 200 nonzero occurrences in the earlier model training sample.'),
'nonzero_training_events':('Training pattern events','Number of eligible nonzero training occurrences for a candlestick pattern. N/A for continuous indicators.'),
'validation_overall':('2025 overall correct','The same selected model\'s overall accuracy in 2025. That year participated in selecting its window profile.'),
'net_return':('Net return after costs','Change in the 10,000 USDT account after every buy/sell fee and slippage. Accuracy is a different question.'),
'buy_hold_return':('Buy and hold','Return from buying once and selling at the end, with the same entry/exit costs.'),
'fee':('Fee each side','0.1% on each buy and each sell in this replay.'),
'slippage':('Slippage each side','Adverse 0.05% price change on each buy and each sell, in addition to fees.'),
'trades':('Completed trades','One buy followed by a sell is one completed trade. Consecutive UP forecasts can stay in one position.'),
'baseline_accuracy':('Simple-guess accuracy','Accuracy of always choosing the direction most common in training, on identical eligible targets.'),
'accuracy_lift':('Lift over simple guess','Overall accuracy minus Simple-guess accuracy. A displayed 2% means two percentage points, not a relative 2% increase.'),
'original_universe_holm_pvalue':('Original adjusted p-value','For original OVERALL accuracy lift only, adjusted across all 410 completed whole-universe comparisons. Does NOT test the new UP-only or DOWN-only percentages. Not reset for this chosen subset.'),
'original_retrospective_gate':('Original accuracy gate','Whether the earlier whole-universe retrospective accuracy gate passed. Not evidence of profit or independent future success.'),
'original_financial_gate':('Original financial gate','Earlier requirement: positive return, beat buy/hold and at least 30 trades in both 2025 and 2026. No earlier model passed.'),
'selected_profile':('Chosen window profile','Default, about 20% shorter or about 20% longer windows, chosen previously using 2025 only. Not changed in this re-test.'),
'parameters':('Indicator settings','The selected numerical formula settings. Window periods count observed candles, not days.'),
'input_series':('Formula inputs','The price/volume series used. CORREL and BETA use BTC high and low series, not a reference asset; market-relative beta is not tested.'),
'source_url':('Source code','The disclosed formula source on GitHub. Native TradingView formula matching is unverified.'),
'status':('Execution status','REEXECUTED_AND_VERIFIED means fresh model execution matched the saved forecasts. On pattern events, NO_SCORED_EVENTS means no measured event accuracy.'),
'scope':('Scoring scope','ALL_ELIGIBLE_CANDLES scores every eligible forecast. CURRENT_CANDLE_PATTERN_FIRED_ONLY restricts the same model to actual pattern occurrences.'),
'input_occurrences':('List appearances','How often you repeated this ID. Repetitions are not separate tests or independent evidence.'),
'evaluation_year':('Walk-forward year','2022, 2023 or 2024. The saved default-profile model was trained only on earlier years.'),
'lift':('Fold accuracy lift','This earlier fold\'s overall accuracy minus its own training-majority baseline, in percentage points.'),
'error':('Recorded problem','The real reason an execution or earlier fit failed. No result is invented.'),
'tp':('Correct UP count','Guessed UP and the actual next close was UP.'),
'tn':('Correct DOWN count','Guessed DOWN and the actual next close was DOWN.'),
'fp':('Wrong UP count','Guessed UP but the actual next close was DOWN.'),
'fn':('Wrong DOWN count','Guessed DOWN but the actual next close was UP.'),
}
FIELDS.update({
 'topic':('Topic','The subject explained on Read Me.'),
 'heading':('Column heading','A header as displayed in the workbook.'),
 'meaning':('Explanation','The plain-language meaning, example or limitation.'),
 'input_row':('Original list row','Position in your 90-row supplied list, starting at 1.'),
 'duplicate_of_row':('Repeated from row','The first earlier row with the same ID. N/A means this is its first appearance.'),
 'path':('Saved data file','Where the verified candle observations are saved inside the repository.'),
 'rows':('Actual candles','The number of observed candles in this dataset, not the number of forecasts or trades.'),
 'first':('First candle opens','Opening time of the first saved candle, in UTC. Indicators use earlier warm-up history before scoring.'),
 'last':('Last candle opens','Opening time of the final saved candle, in UTC. Its closing time is one chart interval later.'),
 'missing_candles':('Missing candle slots','Incomplete or absent observations across the full history. All are before 2025/2026 scoring; no invented candles are filled.'),
 'sha256':('Data fingerprint','The saved price file fingerprint used to check that bytes have not changed.'),
 'derivation':('How prices were built','Official 5m observations or complete aggregations into 15m and 60m candles.'),
 'profile':('Checked window profile','The specific default/short/long profile used for this integrity check.'),
 'kind':('Integrity check type','prefix = regenerate a shorter history; future_mutation = change later prices/volume and check earlier outputs.'),
 'cut':('History bar cutoff','Observation number where the earlier-history sample ends or where later prices begin to change. A bar count, not a date.'),
 'passed':('Integrity check passed','TRUE means earlier numeric outputs stayed unchanged in this particular check. Not a profit pass or proof of complete safety.'),
 'compared_probabilities':('Forecast values compared','Number of original float64 probabilities checked against re-executed predictions, including unscored flat-target rows.'),
 'same_decisions':('Identical UP/DOWN calls','TRUE means every regenerated classification matched the original saved decision.'),
 'maximum_probability_difference':('Largest replay difference','Largest absolute difference between an original and regenerated UP probability. Allowed numerical tolerance is 1e-12; zero is exact reproduction.'),
 'saved_probability_sha256':('Original forecast fingerprint','Fingerprint of the saved probability artifact used for the replay comparison.'),
})
PCTS={'overall_accuracy','up_accuracy','down_accuracy','up_moves_found','down_moves_found','balanced_accuracy','predicted_up_share','actual_up_share','up_ci_low','up_ci_high','down_ci_low','down_ci_high','coverage','validation_overall','net_return','buy_hold_return','fee','slippage','baseline_accuracy','accuracy_lift','lift'}
MAIN=['id','name','overall_accuracy','up_accuracy','up_calls','down_accuracy','down_calls','balanced_accuracy','validation_overall','warnings','pattern_scored_count','net_return','buy_hold_return','source_url']
DETAIL=['id','name','timeframe_minutes','period','scope','status','overall_accuracy','up_accuracy','up_calls','down_accuracy','down_calls','forecasts','up_moves_found','down_moves_found','balanced_accuracy','actual_up_share','predicted_up_share','up_ci_low','up_ci_high','down_ci_low','down_ci_high','coverage','warnings','tp','tn','fp','fn','pattern_fired_count','pattern_scored_count','rare_training_events','nonzero_training_events','net_return','buy_hold_return','trades','fee','slippage','baseline_accuracy','accuracy_lift','original_universe_holm_pvalue','original_retrospective_gate','original_financial_gate']

def warnings(r):
    flags=[]
    for direction in ('up','down'):
        n=r.get(direction+'_calls',0)
        if n==0:flags.append('Never guessed '+direction.upper())
        elif n<100:flags.append('<100 '+direction.upper()+' guesses')
    if r.get('rare_training_events'):flags.append('<200 training pattern events')
    if r.get('pattern_scored_count')==0:flags.append('No scored pattern occurrences')
    elif r.get('pattern_few_scored'):flags.append('<100 scored pattern occurrences')
    if r['id'] in {'BETA','CORREL'}:flags.append('BTC high/low inputs, not another market')
    return '; '.join(flags) or 'No threshold warning; not proof of reliability'

def export():
    verify_inputs();x=json.loads((OUT/'combined.json').read_text());original,ids,counts,c=request()
    rows=[dict(r,warnings=warnings(r)) for r in x['rows']]
    final=[r for r in rows if r['period']=='retrospective']
    for r in final:r['validation_overall']=next(v['overall_accuracy'] for v in rows if v['id']==r['id'] and v['timeframe_minutes']==r['timeframe_minutes'] and v['period']=='validation')
    temp=DEST.with_suffix('.tmp.xlsx');w=xlsxwriter.Workbook(temp,{'strings_to_formulas':False,'strings_to_urls':False})
    head=w.add_format({'bold':True,'bg_color':'#17365D','font_color':'white','text_wrap':True});title=w.add_format({'bold':True,'font_size':15});note=w.add_format({'text_wrap':True,'font_color':'#8B0000','valign':'top'});pct=w.add_format({'num_format':'0.00%'});wrap=w.add_format({'text_wrap':True,'valign':'top'});meta={}
    def table(name,data,keys,explain):
        if set(keys)-set(FIELDS):raise ValueError('Missing header definitions: '+str(set(keys)-set(FIELDS)))
        s=w.add_worksheet(name);s.write(0,0,name,title);s.merge_range(1,0,2,len(keys)-1,explain,note);s.set_row(1,30);s.set_row(3,42)
        for j,k in enumerate(keys):
            s.write(3,j,FIELDS.get(k,(k,''))[0],head);s.set_column(j,j,18)
            if k in FIELDS:s.write_comment(3,j,FIELDS[k][1])
        for i,row in enumerate(data,4):
            for j,key in enumerate(keys):
                v=row.get(key)
                if v is None or isinstance(v,float) and not math.isfinite(v):s.write(i,j,'N/A')
                elif isinstance(v,bool):s.write_boolean(i,j,v)
                elif isinstance(v,(int,float)):s.write_number(i,j,v,pct if key in PCTS else None)
                elif isinstance(v,(list,dict)):s.write_string(i,j,json.dumps(v,sort_keys=True))
                elif key=='source_url':s.write_url(i,j,v)
                else:s.write_string(i,j,str(v))
        s.freeze_panes(4,min(2,len(keys)));s.autofilter(3,0,3+len(data),len(keys)-1)
        for k in ('name','warnings','input_series','error','source_url'):
            if k in keys:s.set_column(keys.index(k),keys.index(k),42,wrap)
        meta[name]={'data_rows':len(data),'keys':keys};return s
    read=[
      dict(topic='What was tested',meaning=f"Your {len(original)} supplied rows contain {len(ids)} different formulas. Completed {len(final)} indicator/timeframe re-executions. Each uses the SAME previously trained logistic model and 0.5 probability threshold."),
      dict(topic='Dates',meaning='Binance spot BTC/USDT. Train 2020–2024; 2025 chose window settings; January–September 2026 is RETROSPECTIVE. These market dates were already examined; this is not independent untouched evidence. All timestamps use UTC.'),
      dict(topic='UP / DOWN accuracy',meaning='Correct UP guesses = correct UP calls divided by all UP calls. Correct DOWN guesses = correct DOWN calls divided by all DOWN calls. N/A means that direction was never predicted, or no events could be scored. Not zero or 100%.'),
      dict(topic='Overall versus recall',meaning='Overall correct counts both kinds of guesses. UP moves found / DOWN moves found instead ask how many ACTUAL moves of each kind were detected. Balanced accuracy averages those two detection rates.'),
      dict(topic='Pattern occurrences',meaning='Main tables score every eligible candle, including times a pattern did not fire. Pattern Events separately scores the SAME model only when the current candlestick pattern actually appeared. Different coverage: do not rank this against all-candle results.'),
      dict(topic='No invented pattern direction',meaning='A positive pattern code is not always a bullish prediction; Doji can simply mark a shape. Direction comes from the frozen model. The model was trained on all eligible candles, not newly fitted only on pattern events.'),
      dict(topic='Volatility is not direction',meaning='ATR, NATR, TRANGE, variance and similar measures are not standalone UP/DOWN forecasts. Results are for indicator + trained readout. CORREL and BETA use BTC high/low series; no other-asset market beta is evaluated. TradingView parity is UNVERIFIED.'),
      dict(topic='Warnings',meaning='Fewer than 100 UP or DOWN calls is flagged. Pattern training counts below 200 and scored pattern events below 100 are flagged. Above these thresholds does not prove reliability. Wilson intervals are descriptive and assume independence; market predictions can be correlated.'),
      dict(topic='Costs',meaning='Accuracy ignores fees. The separate long/cash portfolio starts 10,000 USDT, fills at the next adjacent open and pays 0.1% fee plus 0.05% slippage on EVERY buy AND sell. DOWN means cash, not a short. Buy/hold has identical entry/exit costs. Minimum order sizes, lot rounding and latency are omitted.'),
      dict(topic='Selection and p-values',meaning='You selected this subset after the original historical results existed. This adds selection bias. Original adjusted p-values retain the original 410-test whole-universe correction; none is recalculated to advertise a new independent winner.'),
      dict(topic='How to read',meaning='Start with your chart Results tab. Check BOTH direction percentages AND their guess counts. Read Warnings. Then inspect Pattern Events for candlesticks. Hover a header for its definition, or open Column Guide.'),
    ]
    s=table('Read Me',read,['topic','meaning'],'Plain-language instructions. Accuracy is not profitability; N/A is not zero.');s.set_column(0,0,30);s.set_column(1,1,125,wrap)
    for i in range(4,4+len(read)):s.set_row(i,75)
    for m,label in [(5,'5m Results'),(15,'15m Results'),(60,'1h Results')]:
        data=sorted([r for r in final if r['timeframe_minutes']==m],key=lambda r:r['requested_subset_rank'])
        table(label,data,MAIN,'Jan–Sep 2026 retrospective. Every eligible candle. Correct UP/DOWN guesses are conditional hit rates; inspect counts and warnings. The predictor is indicator + frozen model.')
    table('2025 Validation',[r for r in rows if r['period']=='validation'],['id','name','timeframe_minutes']+[k for k in MAIN if k not in {'id','name','validation_overall'}],'2025 participated in selecting the saved window profile. These are not independent checks of that selection.')
    table('All Results',rows,DETAIL,'Both periods; every eligible candle. Previous source, model and prediction artifacts were verified before replay. N/A values are not invented scores.')
    events=[dict(r,warnings=warnings(r)) for r in x['pattern_events']]
    event_keys=['id','name','timeframe_minutes','period','status','pattern_fired_count','pattern_scored_count','overall_accuracy','up_accuracy','up_calls','down_accuracy','down_calls','up_moves_found','down_moves_found','balanced_accuracy','pattern_flat_targets','pattern_other_unscored','warnings','rare_training_events','nonzero_training_events','source_url']
    table('Pattern Events',events,event_keys,'Only candles where the candlestick pattern fired. SAME trained model, not native bullish/bearish codes. Zero events means no accuracy. Smaller and differing samples prevent a like-for-like all-candle ranking.')
    table('Walk Forward',x['walk_forward'],['id','timeframe_minutes','evaluation_year','status','overall_accuracy','up_accuracy','up_calls','down_accuracy','down_calls','baseline_accuracy','lift','error'],'Original default-profile expanding models re-executed on 2022/2023/2024. Each fit used earlier years. Original failed fits remain failed.')
    table('Source and Settings',final,['id','name','timeframe_minutes','selected_profile','parameters','input_series','input_occurrences','source_url'],'The sources/settings actually used. All native TradingView formula matching is unverified. BETA/CORREL use high/low, not another asset.')
    req=[dict(input_row=i+1,id=n,name=c[n]['name'],input_occurrences=counts[n],duplicate_of_row=(original.index(n)+1 if original.index(n)!=i else None)) for i,n in enumerate(original)]
    table('Request List',req,['input_row','id','name','input_occurrences','duplicate_of_row'],'Every supplied row preserved. Repeated IDs are tested once per timeframe, not counted as independent evidence.')
    data=json.loads((ROOT/'data/predictions/manifest.json').read_text())['datasets']
    table('Data Sources',data,['timeframe_minutes','path','rows','first','last','missing_candles','sha256','derivation'],'Official checksum-verified Binance 5m observations and complete 15m/60m aggregations. No missing validation/final candles. Full 81-archive evidence: data/predictions/manifest.json.')
    failed=x['failures']+[r for r in x['walk_forward'] if r['status']=='ORIGINAL_FOLD_FAILED']
    if not failed:failed=[dict(status='NO_EXECUTION_FAILURES',error='No replay failures. Missing direction calls or pattern events are N/A in their own sheets, not invented results.')]
    table('Failures',failed,['id','timeframe_minutes','evaluation_year','status','error'],'Re-execution errors and previously failed walk-forward fits. A financial loss or no pattern occurrence is different from a software failure.')
    table('Integrity Checks',x['causality'],['id','timeframe_minutes','profile','kind','cut','passed'],'Earlier numerical outputs regenerated with shorter histories or changed future prices/volume. Passing these checks is not proof of every possible bug being absent.')
    table('Reproduction',x['reproduction'],['id','timeframe_minutes','compared_probabilities','same_decisions','maximum_probability_difference','saved_probability_sha256'],'Every regenerated probability and direction compared to the original float64 prediction artifact. This is an actual re-execution, not new parameter tuning.')
    guide=[dict(heading=label,meaning=explanation) for _,(label,explanation) in FIELDS.items()]
    guide += [dict(heading='N/A',meaning='No guess in that direction, no scored pattern events, missing value or undefined statistic. Never assume it means zero.'),dict(heading='Technical evidence tabs',meaning='Rows = actual data observations. SHA256 = file fingerprint. Cut = prefix bar number. Passed = this integrity check passed, not profit. Input row and Duplicate of row refer to your original list. Probability difference compares numerical replay outputs.')]
    g=table('Column Guide',guide,['heading','meaning'],'Definitions for every prediction/performance heading. Exact program fields and complete data artifacts are saved in reports/selected_directional/.');g.set_column(0,0,35);g.set_column(1,1,125,wrap)
    for i in range(4,4+len(guide)):g.set_row(i,45)
    charts=w.add_worksheet('Charts');charts.write(0,0,'Requested subset: overall directional accuracy',title)
    for j,label in enumerate(['5m Results','15m Results','1h Results']):
        chart=w.add_chart({'type':'column'});chart.add_series({'name':label,'categories':[label,4,0,13,0],'values':[label,4,2,13,2]});chart.set_title({'name':label+' — top 10 by overall accuracy'});chart.set_y_axis({'num_format':'0%','min':.45,'max':.55});chart.set_size({'width':920,'height':330});charts.insert_chart(2+j*18,0,chart)
    meta['Charts']={'data_rows':0,'keys':[],'charts':3}
    w.close()
    with zipfile.ZipFile(temp) as z:
        if z.testzip():raise ValueError('Corrupt workbook')
    b=openpyxl.load_workbook(temp,read_only=True,data_only=True)
    for label in ['5m Results','15m Results','1h Results']:
        if b[label].max_row!=56:raise ValueError('Missing requested indicators in '+label)
    b.close();temp.replace(DEST);shutil.copyfile(DEST,OUT/'BTC_UP_DOWN_ACCURACY.xlsx')
    write_json(OUT/'workbook.json',dict(created_at=utc_now(),path=str(DEST.relative_to(ROOT)),bytes=DEST.stat().st_size,sha256=sha256(DEST.read_bytes()),sheets=meta,completed_indicator_timeframes=len(final),unique_indicators=len(ids),input_rows=len(original),results_sha256=sha256((OUT/'combined.json').read_bytes())))
    print('Exported',DEST,DEST.stat().st_size,'bytes')

if __name__=='__main__':export()
