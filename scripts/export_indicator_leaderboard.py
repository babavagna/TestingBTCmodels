#!/usr/bin/env python3
"""Export only executed measurements, keeping missing values and source probes explicit."""
import json,math,sys,zipfile
from pathlib import Path
from datetime import datetime,timezone
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import xlsxwriter
from btc_lab.data import sha256,write_json,utc_now
from indicator_lab.data import ROOT
from indicator_lab.features import catalog,PINE_MATCH
from scripts.indicator_column_guide import guide_rows,write_markdown
OUT=ROOT/'reports/predictions';DEST=ROOT/'reports/TradingView_Indicator_Predictions.xlsx'
PERCENT={'accuracy','validation_accuracy','baseline_majority_accuracy','accuracy_lift_vs_train_majority','lift_block_ci_low','lift_block_ci_high','balanced_accuracy','up_precision','up_recall','actual_up_rate','predicted_up_rate','wilson_low','wilson_high','coverage','net_return','max_drawdown','win_rate','buy_hold_return','excess_vs_buy_hold','fee','slippage','mean_walk_forward_lift','baseline_accuracy','lift'}
MAIN=['timeframe_minutes','retrospective_rank','id','name','group','status','selected_profile','selected_parameters','accuracy','validation_accuracy','baseline_majority_accuracy','accuracy_lift_vs_train_majority','lift_block_ci_low','lift_block_ci_high','holm_pvalue','retrospective_validation_gate_passed','genuine_untouched_oos_status','forecasts','coverage','rare_training_events','nonzero_training_events','too_few_forecasts','net_return','sharpe','max_drawdown','profit_factor','profit_factor_note','win_rate','trades','too_few_trades','buy_hold_return','financial_gate_passed','mean_walk_forward_lift','walk_forward_completed','score_start','score_stop_exclusive','first_forecast_at','last_target_close','data_source','exchange','fee','slippage','github_url','formula_status','reason']
GLOSSARY=[
('Accuracy','If 53 out of 100 next directions were right, accuracy is 53%. It ignores how big the moves were and trading costs.'),
('Forecast target','After this candle closes, predict whether the next CLOSE is higher or lower than this close. Equal closes are excluded from accuracy.'),
('Indicator + readout','An indicator is a calculation, not automatically a prediction. We teach the SAME simple logistic model to read each indicator using past training data.'),
('Balanced accuracy','Average of the hit rate on UP candles and DOWN candles. Helps when one direction happens more often.'),
('Precision / recall','UP precision: how often an UP guess was right. UP recall: how many actual UP moves the model found.'),
('Brier / log loss','Penalize wrong probabilities, especially confident wrong guesses. Lower is better; these are not percentages.'),
('MCC','A balanced correctness score: +1 perfect, 0 no useful association, -1 perfectly reversed. N/A if the denominator is zero.'),
('Baseline','A simple comparison. Training-majority guesses whichever direction happened most in TRAINING, without peeking at the test.'),
('Accuracy lift','Model accuracy minus baseline accuracy. 0.02 means two extra correct answers per 100, or 2 percentage points.'),
('Block confidence interval','A rough uncertainty range using seven-day chunks, preserving some clustering. A lower lift bound <=0 means a benefit is uncertain.'),
('Holm p-value','An approximate statistical check adjusted because we tried many indicators. A small value is evidence, not a promise; it cannot undo earlier research.'),
('Retrospective rank','Ordering by the 2026 accuracy already calculated. The user chose dates examined earlier, so this is NOT a genuinely untouched-test ranking.'),
('Retrospective validation gate','Requires a positive baseline lift in 2025 AND 2026, positive lower 2026 block interval, Holm p<0.05, positive mean earlier walk-forward lift and adequate events/samples. Still retrospective.'),
('Walk forward','Learn only from earlier years, then predict the next year. Default settings: train through 2021/test 2022, through 2022/test 2023, through 2023/test 2024.'),
('Parameter stress','Try windows 20% shorter and longer. 2025 alone picks the final profile. The other 2026 profiles are sensitivity checks, not alternative winners.'),
('Net return','Change in money after fees and slippage. Starting with 10,000 and finishing with 8,000 means -20%.'),
('Sharpe','Average DAILY account return compared with how bumpy daily returns are, scaled by sqrt(365). Higher is usually better; risk-free rate is zero.'),
('Maximum drawdown','Biggest observed-close fall from an earlier account peak. A 40% drawdown means 10,000 temporarily fell to 6,000 after its peak.'),
('Profit factor','Total money won on completed trades divided by money lost, after costs. Below 1 means losses were larger. No losing dollars makes this ratio N/A.'),
('Trade win rate','Share of completed buy/sell trades that made money after costs. This is DIFFERENT from next-candle prediction accuracy.'),
('Trade count','One buy followed by a sell is one completed trade. Consecutive UP guesses can stay in the same trade. Fewer than 30 is flagged.'),
('Buy and hold','Buy Bitcoin once and sell at the end, with the same fees, slippage and start/end times.'),
('Financial gate','Positive net return, above costed buy-and-hold and at least 30 trades in BOTH 2025 and Jan–Sep 2026. Separate from accuracy.'),
('Slippage','Assume every purchase happens at 0.05% above the displayed price, and every sale 0.05% below. Add the 0.1% fee on each side.'),
('Too few events','A candlestick pattern with fewer than 200 nonzero training occurrences is flagged. A rare pattern can appear accurate mainly by guessing the usual direction.'),
('N/A','Missing, not executed, unverified or mathematically undefined. It NEVER means zero.'),
('TradingView parity','UNVERIFIED for every formula. Open-source formulas, defaults and initialization can differ from the native TradingView engine.'),
('Scope limit','147 reproducible formula entries were attempted per timeframe, including 61 candlestick patterns. Separately audit 210 public Pine scripts. This is NOT all TradingView indicators.'),
]

def main():
    combined=json.loads((OUT/'combined.json').read_text());rows=combined['rows'];final=[r.copy() for r in rows if r['period']=='retrospective']
    manifest=json.loads((ROOT/'data/predictions/manifest.json').read_text());audit=json.loads((ROOT/'sources/tradingview/pine_audit.json').read_text())
    freezes=[json.loads((OUT/f'frozen_{m}m.json').read_text()) for m in (5,15,60)]
    for r in final:
        val=next(v for v in rows if v['id']==r['id'] and v['timeframe_minutes']==r['timeframe_minutes'] and v['period']=='validation')
        r.update(validation_accuracy=val['accuracy'],data_source='Binance Vision official monthly BTCUSDT 5m archives; complete 15m/60m aggregation',exchange='Binance spot BTC/USDT',reason='')
    final.sort(key=lambda r:(r['timeframe_minutes'],r['retrospective_rank']))
    allrows=[]
    for minutes in (5,15,60):
        for item in catalog():
            found=next((r for r in final if r['id']==item['id'] and r['timeframe_minutes']==minutes),None)
            if found:allrows.append(found)
            else:
                reasons=[f['error'] for z in freezes for f in z['failures'] if f['id']==item['id'] and f['timeframe_minutes']==minutes]
                reasons += [f['error'] for f in combined['score_failures'] if f['id']==item['id'] and f['timeframe_minutes']==minutes]
                allrows.append(dict(**item,timeframe_minutes=minutes,status='FAILED_NO_FORECAST_TEST',reason='; '.join(sorted(set(reasons))),genuine_untouched_oos_status='UNAVAILABLE',data_source=manifest['source'],exchange='Binance spot BTC/USDT',fee=.001,slippage=.0005))
    failures=[r for f in freezes for r in f['failures']]+combined['score_failures']
    failures += [dict(id=w['id'],timeframe_minutes=f['timeframe_minutes'],stage=f"walk_forward_{w['evaluation_year']}",error=w['error'],forecast_test_completed=False) for f in freezes for w in f['walk_forward'] if w['status']=='failed']
    failures += [dict(id=p['id'],stage='ORIGINAL_PINE_COMPATIBILITY_ONLY',error=p.get('error'),forecast_test_completed=False,github_url=p['github_url']) for p in audit['scripts'] if p['status']!='compatibility_probe_completed_only']
    temporary=DEST.with_name(DEST.stem+'.tmp.xlsx')
    w=xlsxwriter.Workbook(temporary,{'strings_to_formulas':False,'strings_to_urls':False,'nan_inf_to_errors':False})
    w.set_properties({'title':'BTC indicator next-candle predictions — retrospective study','comments':'Executed formula benchmark, not all TradingView scripts. No native TradingView parity claim.'})
    fmt={
      'title':w.add_format({'bold':True,'font_size':15,'font_color':'#17365D'}),
      'note':w.add_format({'font_color':'#8B0000','text_wrap':True,'valign':'top'}),
      'head':w.add_format({'bold':True,'bg_color':'#17365D','font_color':'white','text_wrap':True,'valign':'top'}),
      'pct':w.add_format({'num_format':'0.00%'}),'num':w.add_format({'num_format':'0.0000'}),
      'text':w.add_format({'text_wrap':True,'valign':'top'}),
      'bad':w.add_format({'bg_color':'#FCE4D6'}),'good':w.add_format({'bg_color':'#E2F0D9'}),
    }
    sheet_meta={}
    def table(name,data,columns=None,note='Executed results only. N/A is missing/undefined. Rankings are RETROSPECTIVE; native TradingView parity UNVERIFIED.'):
        sh=w.get_worksheet_by_name(name) or w.add_worksheet(name);cols=columns or list(dict.fromkeys(k for r in data for k in r))
        if not cols:cols=['status'];data=[{'status':'No failures recorded'}]
        sh.write(0,0,name,fmt['title']);sh.merge_range(1,0,2,max(1,len(cols)-1),note,fmt['note']);sh.set_row(1,24);sh.set_row(3,44)
        for j,k in enumerate(cols):sh.write(3,j,k,fmt['head']);sh.set_column(j,j,18 if k not in {'name','github_url','error','reason','formula_status'} else 40)
        for i,r in enumerate(data,4):
            for j,k in enumerate(cols):
                value=r.get(k)
                if value is None or (isinstance(value,float) and not math.isfinite(value)):sh.write(i,j,'N/A');continue
                if isinstance(value,(dict,list)):value=json.dumps(value,sort_keys=True)
                if isinstance(value,bool):sh.write_boolean(i,j,value)
                elif isinstance(value,(float,int)):sh.write_number(i,j,value,fmt['pct'] if k in PERCENT else (fmt['num'] if isinstance(value,float) else None))
                elif k in {'github_url','url','checksum_url'} and str(value).startswith('https://'):sh.write_url(i,j,value,string=value)
                else:sh.write_string(i,j,str(value))
        sh.freeze_panes(4,min(3,len(cols)));sh.autofilter(3,0,3+len(data),len(cols)-1)
        for k in ['net_return','accuracy_lift_vs_train_majority','excess_vs_buy_hold']:
            if k in cols and data:
                col=cols.index(k);sh.conditional_format(4,col,3+len(data),col,{'type':'cell','criteria':'<','value':0,'format':fmt['bad']})
        sheet_meta[name]={'rows':len(data),'columns':cols};return sh
    summary=[dict(item='Study',explanation='BTCUSDT next-close prediction at 5m, 15m and 60m. Indicator + trained logistic model, not an indicator magically forecasting by itself.'),dict(item='Execution coverage',explanation=f"{len(final)} completed indicator/timeframe forecast tests out of 441 attempts. See All Strategies and Failed Tests. 210 ORIGINAL Pine probes are separate and do not count as forecast tests."),dict(item='Test dates',explanation='Training 2020–2024; validation 2025; RETROSPECTIVE test January–September 2026. Those market dates were already examined. NO genuinely untouched OOS result is available.'),dict(item='Assumptions',explanation='Binance spot BTC/USDT; 10,000 USDT; next-open long/cash fills; fee 0.1% + slippage 0.05% on EVERY buy AND sell; final close exit. Accuracy is separate from trading profitability.'),dict(item='How to read',explanation='Top 10 has ten retrospective rows for EACH timeframe. Look at baseline lift, validation, sample flags, costs and robustness before interpreting a high hit rate.'),dict(item='Download',explanation='Save this .xlsx locally and open with Excel or LibreOffice. GitHub may not preview it. Use the direct raw download link in README.'),dict(item='Reproduce',explanation='Read indicator_lab/README.md. Frozen models, exact probabilities/decisions, labels, prices, source bodies, hashes and code are saved in this repository.')]
    summary += [dict(item=k,explanation=v) for k,v in GLOSSARY]
    summary.insert(1,dict(item='Every column explained',explanation='Open the Column Guide tab beside Read Me. It explains every heading, examples, TRUE/FALSE warnings and N/A values across all sheets. Accuracy means correct directions; net return means money after costs.'))
    next(r for r in summary if r['item']=='Download')['explanation']='Open DOWNLOAD_BTC_INDICATOR_RESULTS.xlsx in the repository top folder while signed in to GitHub and choose Download raw file. If that fails, Code > Download ZIP downloads the whole project. Save locally and open with Excel or LibreOffice.'
    read=table('Read Me',summary,['item','explanation']);read.set_column(0,0,30);read.set_column(1,1,125)
    for i in range(4,4+len(summary)):read.set_row(i,50)
    w.add_worksheet('Column Guide')
    table('All Strategies',allrows,MAIN)
    table('Top 10',[r for r in final if r['retrospective_rank']<=10],MAIN)
    for minutes in (5,15,60):table(f'{minutes}m Accuracy',[r for r in final if r['timeframe_minutes']==minutes],MAIN)
    accuracy_cols=['timeframe_minutes','id','name','period','accuracy','balanced_accuracy','up_precision','up_recall','mcc','brier','log_loss','actual_up_rate','predicted_up_rate','forecasts','coverage','tp','tn','fp','fn','wilson_low','wilson_high','baseline_majority_accuracy','accuracy_lift_vs_train_majority','lift_block_ci_low','lift_block_ci_high','holm_pvalue','retrospective_validation_gate_passed','too_few_forecasts','rare_training_events']
    table('Prediction Accuracy',rows,accuracy_cols)
    table('Baselines',combined['baselines'])
    wf=[{k:v for k,v in r.items() if k!='model'} for f in freezes for r in f['walk_forward']]
    table('Walk Forward',wf)
    table('Robustness',combined['robustness'])
    table('Monthly Accuracy',combined['monthly'])
    table('Failed Tests',failures,['id','timeframe_minutes','profile','stage','error','forecast_test_completed','github_url'])
    table('Pine Audit',audit['scripts'],['id','name','pine_version','status','forecast_status','native_tradingview_parity','risk_flags','error','plots','github_url','sha256'],note='Compatibility probes of unchanged original scripts on 1,000 candles. Even a successful probe is NOT a forecast test. Related Python formula tests appear separately.')
    table('Sources',catalog())
    table('Price Data',manifest['datasets'])
    table('Monthly Archives',manifest['months'])
    checks=[dict(timeframe_minutes=f['timeframe_minutes'],**r) for f in freezes for r in f['causality_checks']]
    table('Causality Checks',checks)
    # Chart data are the same scored values; no smoothed or invented chart series.
    cdata=[]
    for r in final:
        if r['retrospective_rank']<=10:cdata.append({k:r[k] for k in ['timeframe_minutes','id','accuracy','baseline_majority_accuracy','net_return','buy_hold_return']})
    table('Chart Data',cdata)
    charts=w.add_worksheet('Charts');charts.write(0,0,'Retrospective accuracy and net return',fmt['title']);charts.write(1,0,'Accuracy ignores costs; every displayed indicator is an executed formula/model test.',fmt['note']);charts.set_column(0,14,12)
    for index,minutes in enumerate((5,15,60)):
        first=4+index*10;last=first+9
        chart=w.add_chart({'type':'column'})
        for col,label in [(2,'Next-close accuracy'),(3,'Training-majority baseline')]:chart.add_series({'name':label,'categories':['Chart Data',first,1,last,1],'values':['Chart Data',first,col,last,col]})
        chart.set_title({'name':f'{minutes}m — 2026 retrospective top 10'});chart.set_y_axis({'num_format':'0%','min':.45,'max':.60});chart.set_legend({'position':'bottom'});chart.set_size({'width':900,'height':330});charts.insert_chart(index*18+3,0,chart)
    chart=w.add_chart({'type':'scatter','subtype':'straight_with_markers'})
    for index,minutes in enumerate((5,15,60)):
        first=4+index*10;last=first+9;chart.add_series({'name':f'{minutes}m','categories':['Chart Data',first,2,last,2],'values':['Chart Data',first,4,last,4],'line':{'none':True},'marker':{'type':'circle','size':5}})
    chart.set_title({'name':'High hit rate can still lose money after costs'});chart.set_x_axis({'name':'Direction accuracy','num_format':'0%'});chart.set_y_axis({'name':'Net return after costs','num_format':'0%'});chart.set_size({'width':900,'height':400});charts.insert_chart(58,0,chart)
    sheet_meta['Charts']={'rows':0,'columns':[],'charts':4}
    definitions=guide_rows(sheet_meta)
    guide=table('Column Guide',definitions,['column','meaning','how_to_read','sheet'],note='Every column explained. TRUE can mean a passed check OR a warning; read its definition. N/A is not zero. Accuracy is not trading profit. Dates use UTC.')
    guide.set_column(0,0,39,fmt['text']);guide.set_column(1,1,60,fmt['text']);guide.set_column(2,2,85,fmt['text']);guide.set_column(3,3,48,fmt['text'])
    guide.freeze_panes(4,1)
    for i in range(4,4+len(definitions)):guide.set_row(i,100)
    w.close()
    with zipfile.ZipFile(temporary) as z:
        if z.testzip() is not None:raise ValueError('Workbook zip integrity failed')
        if len([n for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')])!=4:raise ValueError('Workbook chart coverage incomplete')
    import openpyxl
    check=openpyxl.load_workbook(temporary,read_only=True,data_only=True)
    if check['All Strategies'].max_row!=445:raise ValueError('Workbook inventory incomplete')
    if check['Column Guide'].max_row!=4+len(definitions):raise ValueError('Column guide incomplete')
    check.close()
    temporary.replace(DEST)
    # A visible top-level copy lets beginners navigate from their signed-in repo.
    import shutil
    shutil.copyfile(DEST,ROOT/'DOWNLOAD_BTC_INDICATOR_RESULTS.xlsx')
    write_markdown(ROOT/'WORKBOOK_COLUMN_GUIDE.md',definitions)
    write_json(DEST.with_suffix('.json'),dict(created_at=utc_now(),path=str(DEST.relative_to(ROOT)),bytes=DEST.stat().st_size,sha256=sha256(DEST.read_bytes()),completed_forecast_tests=len(final),attempted_indicator_timeframe_tests=441,original_pine_compatibility_probes=210,holm_family_size=combined['holm_family_size'],ranking='Retrospective, not genuinely untouched',sheets=sheet_meta,inputs={str(p.relative_to(ROOT)):sha256(p.read_bytes()) for p in [OUT/'combined.json',ROOT/'data/predictions/manifest.json',ROOT/'sources/tradingview/pine_audit.json']},automated_test_evidence='reports/predictions/verification.json'))
    print(DEST,DEST.stat().st_size,'bytes',len(final),'completed forecast tests')
if __name__=='__main__':main()
