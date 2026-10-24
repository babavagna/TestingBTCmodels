"""Measured requested-list results must trace to actual frozen predictions and prices."""
import json,zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
import openpyxl
from btc_lab.data import sha256
from indicator_lab.data import ROOT,load
from indicator_lab.features import raw
from directional_study.run import request,verify_inputs
from directional_study.metrics import directions
OUT=ROOT/'reports/selected_directional'
pytestmark=pytest.mark.skipif(not (OUT/'combined.json').exists(),reason='Requested-list study has not executed')

def result():return json.loads((OUT/'combined.json').read_text())

def test_all_156_requested_model_replays_completed_and_match_frozen_evidence():
    verify_inputs();x=result();assert len(x['rows'])==312 and not x['failures']
    assert len(x['reproduction'])==156 and all(r['same_decisions'] and r['maximum_probability_difference']<=1e-12 for r in x['reproduction'])
    assert len(x['causality'])==624 and all(r['passed'] for r in x['causality'])
    assert len({(r['id'],r['timeframe_minutes']) for r in x['rows']})==156

def test_every_up_down_score_and_pattern_only_score_recomputes_from_true_targets():
    x=result();_,ids,_,catalog=request()
    fields=['forecasts','up_accuracy','down_accuracy','up_calls','down_calls','overall_accuracy','up_moves_found','down_moves_found','tp','tn','fp','fn']
    for m in (5,15,60):
        d,_=load(m);labels=pd.read_csv(ROOT/f'reports/predictions/labels_{m}m.csv.gz',parse_dates=['feature_candle_open','forecast_available_at','target_close_at'])
        positions=d.index.get_indexer(labels.feature_candle_open);assert (positions>=0).all()
        for ident in ids:
            with np.load(ROOT/f'reports/predictions/probabilities/{ident}_{m}m.npz') as a:p=a['probability'];good=a['eligible_features']
            for row in [r for r in x['rows'] if r['id']==ident and r['timeframe_minutes']==m]:
                start=pd.Timestamp(row['score_start'],tz='UTC');end=pd.Timestamp(row['score_stop_exclusive'],tz='UTC')
                mask=(labels.forecast_available_at>=start)&(labels.target_close_at<=end)&labels.scorable_direction&good
                a=directions(labels.actual_up[mask].to_numpy(),p[mask])
                assert all(a[k]==row[k] for k in fields)
                if row['pattern']:
                    output=raw(catalog[ident],d,row['selected_profile'])
                    active=np.isfinite(output).all(axis=1)&np.any(output!=0,axis=1)
                    eventmask=mask&active[positions]
                    ev=next(e for e in x['pattern_events'] if e['id']==ident and e['timeframe_minutes']==m and e['period']==row['period'])
                    a=directions(labels.actual_up[eventmask].to_numpy(),p[eventmask])
                    assert all(a[k]==ev[k] for k in fields)
                    assert ev['pattern_scored_count']==row['pattern_scored_count']==eventmask.sum()
                    assert ev['pattern_fired_count']==ev['pattern_scored_count']+ev['pattern_flat_targets']+ev['pattern_other_unscored']

def test_no_direction_calls_or_no_events_never_receive_a_fabricated_accuracy():
    x=result()
    for r in x['rows']+x['pattern_events']:
        assert r['up_calls']+r['down_calls']==r['forecasts']
        for direction in ['up','down']:
            if r[direction+'_calls']==0:assert r[direction+'_accuracy'] is None
        if r['forecasts']==0:assert r['overall_accuracy'] is None
    assert any(r['forecasts']==0 for r in x['pattern_events'])
    assert any(r['down_calls']==0 for r in x['rows'])

def test_the_selected_subset_does_not_reset_the_original_significance_family():
    old=json.loads((ROOT/'reports/predictions/combined.json').read_text())
    for row in [r for r in result()['rows'] if r['period']=='retrospective']:
        original=next(r for r in old['rows'] if r['id']==row['id'] and r['timeframe_minutes']==row['timeframe_minutes'] and r['period']=='retrospective')
        assert row['original_universe_holm_pvalue']==original['holm_pvalue']
        assert row['original_retrospective_gate']==original['retrospective_validation_gate_passed']
        assert row['original_financial_gate']==original['financial_gate_passed']
        assert row['net_return']==pytest.approx(original['net_return'],abs=1e-12)

def test_focused_workbook_is_valid_complete_and_keeps_na_directions():
    p=ROOT/'BTC_UP_DOWN_ACCURACY.xlsx';meta=json.loads((OUT/'workbook.json').read_text())
    assert sha256(p.read_bytes())==meta['sha256'] and p.read_bytes()==(OUT/p.name).read_bytes()
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None
        assert len([n for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')])==3
    b=openpyxl.load_workbook(p,read_only=True,data_only=True)
    for name in ['5m Results','15m Results','1h Results']:
        s=b[name];assert s.max_row==56
        head=[c.value for c in s[4]]
        for r in s.iter_rows(min_row=5,values_only=True):
            if r[head.index('DOWN guesses')]==0:assert r[head.index('Correct DOWN guesses')]=='N/A'
    assert b['Request List'].max_row==94
    assert b['Pattern Events'].max_row==214
    guide={r[0] for r in b['Column Guide'].iter_rows(min_row=5,values_only=True)}
    for s in b:
        if s.title!='Charts':assert {c.value for c in s[4] if c.value}<=guide
    b.close()
