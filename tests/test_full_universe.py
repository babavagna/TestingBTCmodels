"""Integrity, causal counterexamples and source-rule checks for the expanded list."""
import importlib.util,json,sys,types,zipfile
from collections import namedtuple
import numpy as np
import pandas as pd
import pytest
import talib
from openpyxl import load_workbook
from btc_lab.data import DataError,load_verified,sha256
from full_universe.audit import catalog,read_report,ROOT,OUT,period_run
from full_universe.jesse_adapter import generate,defaults
from full_universe.source_review import indicator_proof


def test_full_catalog_contains_exact_supplied_sources():
    entries=catalog();cross=pd.read_csv(ROOT/'sources/candidate_crosswalk.csv')
    assert len(entries)==100
    assert {i['framework'] for i in entries}=={'freqtrade','jesse','zenbot'}
    for i in entries:
        original=cross.loc[cross.list_number==i['number']].iloc[0]
        assert i['source_url']==original.pinned_github_url
        assert sha256((ROOT/i['source_path']).read_bytes())==i['source_sha256']
        assert i['native_execution_reproduced'] is False


def test_100_attempts_never_present_failed_candidates_as_backtested():
    status,periods,meta=read_report()
    assert len(status)==100 and status.strategy_id.nunique()==100
    completed=set(status.loc[status.status=='completed_signal_backtest','strategy_id'])
    assert set(periods.strategy_id)==completed
    assert len(periods)==len(completed)*2
    assert periods.execution_status.eq('completed').all()
    assert not status.unseen_rank_eligible.any()
    assert meta['untouched_ranks']==0 and meta['directional_models_scored']==0
    assert periods.groupby('strategy_id').period.nunique().eq(2).all()
    assert not set(status.loc[status.status!='completed_signal_backtest','strategy_id'])&set(periods.strategy_id)


def test_selected_parameters_and_trades_respect_purged_date_boundaries():
    _,periods,_=read_report()
    frozen=json.loads((OUT/'frozen_selection.json').read_text())
    assert all(pd.Timestamp(r['last_fit_candle'])==pd.Timestamp('2024-12-31T22:00Z') for r in frozen)
    for row in periods.itertuples():
        trades=pd.read_csv(OUT/'candidates'/f'{row.strategy_id}_{row.period}_trades.csv')
        if len(trades):
            start=pd.Timestamp(row.start,tz='UTC');stop=pd.Timestamp(row.stop_exclusive,tz='UTC')
            assert pd.to_datetime(trades.entry_time,utc=True).ge(start).all()
            assert pd.to_datetime(trades.exit_time,utc=True).le(stop).all()
        assert len(trades)==row.trades
        pnl=trades.net_pnl_usdt
        assert row.return_pct==pytest.approx(pnl.sum()/10000*100,abs=1e-8)


def test_detects_actual_indicator_leak_even_without_signal_threshold_crossing():
    data,_=load_verified(ROOT)
    for number in (24,36):
        proof=indicator_proof(catalog()[number-1],data)
        assert proof['changed_cells']>0 and proof['examples']
        assert all(pd.Timestamp(e['timestamp'])<pd.Timestamp(proof['cutoff']) for e in proof['examples'])
        r=json.loads((OUT/'candidates'/f's{number:03}.json').read_text())
        assert r['status']['status']=='quarantined_lookahead'
        assert not r['period_results']
        assert not list((OUT/'candidates').glob(f's{number:03}_*_trades.csv'))


@pytest.mark.parametrize('number',[53,62])
def test_unsupported_native_state_fails_instead_of_substituting_another_strategy(number,candles):
    with pytest.raises(NotImplementedError,match='UNSUPPORTED_NATIVE_STATE'):
        generate(catalog()[number-1],candles)


def test_fresh_cash_benchmark_costs_and_no_future_trade_crossing():
    data,_=load_verified(ROOT)
    stats,eq,trades=period_run(data,None,'2026-01-01','2026-10-01')
    expected=((data.loc['2026-09-30T23:00Z','close']*(1-.0005)*(1-.001))/(data.loc['2026-01-01T00:00Z','open']*(1+.0005)*(1+.001))-1)*100
    assert stats['return_pct']==pytest.approx(expected)
    assert eq.index[0]==pd.Timestamp('2026-01-01T00:00Z')
    assert len(trades)==1


@pytest.mark.parametrize('number',[54,55,56,57,58,59,60,61])
def test_jesse_translation_matches_actual_source_rules_with_documented_indicator_backend(number,candles,monkeypatch):
    # Execute the original methods with TA-Lib/pandas indicators. This validates
    # the rule translation, not equivalence to every Jesse indicator version.
    item=catalog()[number-1];d=candles
    modules={k:types.ModuleType(k) for k in ('jesse','jesse.strategies','jesse.indicators','jesse.utils')}
    class Strategy:
        def __init__(self):self.vars={};self.is_long=True;self.is_short=False;self.balance=10000;self.available_margin=1000000;self.fee_rate=.001;self.exited=False
        def liquidate(self):self.exited=True
    modules['jesse.strategies'].Strategy=Strategy;modules['jesse.strategies'].cached=lambda fn:fn
    ta=modules['jesse.indicators'];u=modules['jesse.utils']
    def source(c,source_type='close'):
        return c[:,{'open':1,'close':2,'high':3,'low':4,'volume':5}[source_type]] if source_type!='hl2' else (c[:,3]+c[:,4])/2
    def scalar(fn):
        def calc(c,period=14,source_type='close',sequential=False,**kwargs):
            result=fn(source(c,source_type),timeperiod=period);return result if sequential else result[-1]
        return calc
    ta.sma=scalar(talib.SMA);ta.ema=scalar(talib.EMA);ta.rsi=scalar(talib.RSI)
    ta.ma=lambda c,matype,period,source_type,sequential:talib.LINEARREG(source(c,source_type),timeperiod=period)
    ta.adx=lambda c,period=14:talib.ADX(c[:,3],c[:,4],c[:,2],timeperiod=period)[-1]
    ta.ht_trendmode=lambda c:talib.HT_TRENDMODE(c[:,2])[-1]
    ta.macd=lambda c,fastperiod=12,slowperiod=26,signalperiod=9:tuple(a[-1] for a in talib.MACD(c[:,2],fastperiod,slowperiod,signalperiod))
    DC=namedtuple('DC','upperband middleband lowerband')
    ta.donchian=lambda c,period=20:DC(np.max(c[-period:,3]),(np.max(c[-period:,3])+np.min(c[-period:,4]))/2,np.min(c[-period:,4]))
    ta.bollinger_bands=lambda c,source_type='close':tuple(a[-1] for a in talib.BBANDS(source(c,source_type),timeperiod=20,nbdevup=2,nbdevdn=2))
    Cloud=namedtuple('Cloud','span_a span_b')
    def cloud(c,conversion_line_period=9,base_line_period=26,lagging_line_period=52,displacement=26):
        h=pd.Series(c[:,3]);l=pd.Series(c[:,4]);mid=lambda p:(h.rolling(p).max()+l.rolling(p).min())/2
        return Cloud(((mid(conversion_line_period)+mid(base_line_period))/2).shift(displacement).iloc[-1],mid(lagging_line_period).shift(displacement).iloc[-1])
    ta.ichimoku_cloud=cloud
    u.size_to_qty=lambda balance,price,*a,**kw:balance/price*.99
    def crossed(a,b,direction):
        return (a[-1]>b and a[-2]<=b) if direction=='above' else (a[-1]<b and a[-2]>=b)
    u.crossed=crossed
    modules['jesse'].utils=u;modules['jesse'].indicators=ta
    for name,mod in modules.items():monkeypatch.setitem(sys.modules,name,mod)
    spec=importlib.util.spec_from_file_location('jesse_original_'+item['id'],ROOT/item['source_path']);mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    cls=getattr(mod,item['name']);obj=cls()
    obj.hp={r['name']:r['default'] for r in obj.hyperparameters()} if hasattr(obj,'hyperparameters') else {}
    en,ex=generate(item,d)
    for k in range(300,len(d),37):
        prefix=d.iloc[:k+1];obj.candles=np.column_stack((prefix.index.asi8/1e6,prefix[['open','close','high','low','volume']].to_numpy()))
        obj.price=obj.close=float(prefix.close.iloc[-1]);obj.exited=False
        buy=bool(obj.should_long()) and all(f() for f in obj.filters()) if hasattr(obj,'filters') else bool(obj.should_long())
        obj.update_position()
        assert bool(en.iloc[k])==buy
        assert bool(ex.iloc[k])==obj.exited


def test_full_excel_has_all_candidates_and_exact_saved_metrics(tmp_path):
    from scripts.export_full_leaderboard import export
    path=tmp_path/'full.xlsx';metadata=export(path);wb=load_workbook(path,data_only=True)
    assert zipfile.ZipFile(path).testzip() is None
    assert metadata['all_candidates']==100
    ws=wb['All Strategies'];headers=[c.value for c in ws[8]]
    rows=[dict(zip(headers,r)) for r in ws.iter_rows(min_row=9,values_only=True) if r[0] is not None]
    assert len(rows)==100
    saved=pd.read_csv(OUT/'period_results.csv').query("period=='retrospective_test_2026'").set_index('strategy')
    for r in rows:
        assert r['Unseen rank']=='NOT ISSUED'
        if r['Execution status']=='completed_signal_backtest':
            assert r['Net return']==pytest.approx(saved.loc[r['Strategy'],'return_pct']/100)
            assert r['Completed trades']==saved.loc[r['Strategy'],'trades']
        else:assert str(r['Net return']).startswith('N/A')
    pred=wb['Prediction Accuracy'];assert pred.max_row==108
    h=[c.value for c in pred[8]];neural=dict(zip(h,[c.value for c in pred[84]]))
    assert neural['Strategy']=='neural'
    assert neural['Provides forecasts?'].startswith('YES')
    assert neural['Forecasts scored']==0
    assert len(wb['Charts']._charts)==4


def test_expanded_gap_policy_uses_only_observed_rows_and_retains_history(monkeypatch):
    from full_universe import adapters
    idx=pd.date_range('2020-01-01',periods=4200,freq='h',tz='UTC').delete(4150)
    d=pd.DataFrame({'close':np.ones(len(idx))},index=idx)
    calls=[]
    def real_rows_only(item,frame,factor):
        calls.append(frame.index)
        return pd.Series(True,index=frame.index),pd.Series(False,index=frame.index)
    monkeypatch.setattr(adapters,'generate',real_rows_only)
    en,ex=adapters.prepared({'framework':'test'},d)
    assert len(calls)==1 and calls[0].equals(idx)
    assert not en.iloc[:adapters.WARMUP-1].any()
    assert en.iloc[adapters.WARMUP-1:].all()
    assert en.index.equals(idx)
    assert not ex.any()


def test_dashboard_shows_full_inventory():
    from streamlit.testing.v1 import AppTest
    app=AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
    assert not app.exception
    assert any(len(t.value)==100 and 'status' in t.value.columns for t in app.dataframe)
    assert any('All 100 candidates attempted' in t.value for t in app.success)


def test_stress_summary_preserves_actual_scenario_coverage_and_walk_forward_evidence():
    read_report()
    summary=pd.read_csv(OUT/'robustness_summary.csv')
    scenarios=pd.read_csv(OUT/'sensitivity.csv')
    walk=pd.read_csv(OUT/'walk_forward.csv')
    for r in summary.itertuples():
        cases=scenarios[(scenarios.strategy_id==r.strategy_id)&(scenarios.period==r.period)]
        assert r.completed_scenarios==len(cases)
        assert r.profitable_scenarios==sum(cases.return_pct>0)
        assert 0<=r.profitable_fraction<=1
        assert r.lowest_return_pct<=r.highest_return_pct
        assert r.selected_double_cost_return_pct==pytest.approx(cases.loc[cases.cost_case=='double_both','return_pct'].iloc[0])
        folds=walk[walk.strategy_id==r.strategy_id]
        assert len(folds)==3
        assert r.profitable_walk_forward_folds==sum(folds.return_pct>0)
        assert r.beat_hold_walk_forward_folds==sum(folds.return_pct>folds.buy_and_hold_return_pct)
