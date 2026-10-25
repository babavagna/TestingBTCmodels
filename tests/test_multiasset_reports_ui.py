"""Verify real saved evidence, spreadsheet contents, and all contract tabs."""
import json
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
import pytest
from btc_lab.data import sha256
from multiasset.registry import ROOT,assets
from multiasset_report import build,case_status,case_records,run_directory


def test_real_checkpoints_have_identical_primary_inputs_and_honest_statuses():
    total=0
    for spec in assets():
        for minutes in [5,15,60]:
            for basis in ['trades','quotes']:
                status=case_status(ROOT,spec['id'],minutes,basis)
                records=case_records(status)
                if not records:continue
                assert 0<len(records)<=247
                if status['status']=='COMPLETED_WITH_EXPLICIT_LIMITATIONS':assert len(records)==247
                else:assert status['status'] in ['RUNNING','PARTIAL_CANDIDATE_CHECKPOINT']
                meta=status['dataset'];assert sha256((ROOT/meta['path']).read_bytes())==meta['sha256']
                assert meta['provider']=='Kalshi'
                stages=list(status['dates'].values())
                assert all(pd.Timestamp(a['stop'])<=pd.Timestamp(b['start']) for a,b in zip(stages,stages[1:]))
                for r in records:
                    assert r['identity']==status['identity']
                    if not r.get('executed'):
                        assert not r.get('periods');continue
                    total+=1
                    assert all(c.get('passed') for c in r['causality_checks'])
                    for period in ['validation','retrospective_test']:
                        block=r['periods'][period]
                        if basis=='quotes':assert 'stats' not in block
                        else:
                            s=block['stats'];assert s['fee']==.001 and s['slippage']==.0005
                            assert s['rank_eligible'] is False and s['funding_included'] is False
                        d=block.get('direction')
                        if d:
                            assert d['forecasts']==sum(d[k] for k in ['tp','tn','fp','fn'])
                            assert d['accuracy']==(d['tp']+d['tn'])/d['forecasts']
                directory=run_directory(status)
                if not (directory/'probabilities.json').exists():continue
                for m in json.loads((directory/'probabilities.json').read_text()):
                    assert m.get('current_probabilities') is None
                    if m['status']=='FITTED_AND_CALIBRATED':
                        assert np.isclose(sum(m['historical_snapshot_probabilities']),1)
                        assert pd.Timestamp(m['training_end'])<=pd.Timestamp(m['dates']['calibration']['start'])
                        assert pd.Timestamp(m['calibration_end'])<=pd.Timestamp(m['dates']['validation']['start'])
                        assert m['model']['calibration_log_loss_after']<=m['model']['calibration_log_loss_before']+1e-10
    assert total>1000


def test_workbook_contains_actual_sortable_numeric_rows_not_empty_tables():
    manifest=json.loads((ROOT/'reports/multiasset/workbook_manifest.json').read_text())
    file=ROOT/manifest['workbook'];assert sha256(file.read_bytes())==manifest['workbook_sha256']
    book=openpyxl.load_workbook(file,read_only=False)
    for name,count in manifest['sheet_rows'].items():
        sheet=book[name]
        assert sheet.max_row==count+5
        assert sheet.tables, name
    assert manifest['sheet_rows']['Coverage']==39*3
    headers={v.value:v.column for v in book['Prediction Accuracy'][5]}
    row=book['Prediction Accuracy'][6]
    assert isinstance(row[headers['directional_accuracy']-1].value,float)
    assert isinstance(row[headers['up_precision']-1].value,(int,float))
    assert isinstance(row[headers['down_precision']-1].value,(int,float,str))
    assert len(book['Charts']._charts)>0
    assert book['Top 10']['A6'].value=='NO FULLY VALIDATED FINANCIAL RANKS'
    assert all(cell.data_type!='f' for sheet in book for line in sheet for cell in line)
    book.close()


def test_every_asset_tab_and_timeframe_opens_and_keeps_original_features():
    from streamlit.testing.v1 import AppTest
    app=AppTest.from_file(str(ROOT/'app.py')).run(timeout=45)
    assert not app.exception
    assert [t.label for t in app.tabs]==[s['id'] for s in assets()]
    assert any('Full 100-candidate comparison' in s.value for s in app.subheader)
    for minutes in [15,60,5]:
        app.selectbox(key='multi_tf').set_value(minutes).run(timeout=45)
        assert not app.exception
        assert app.selectbox(key='multi_tf').value==minutes
    assert any('No current probability' in m.value for m in app.info) or any(m.label=='UP' for m in app.metric)
    app.selectbox(key='replay_pick_BTC').set_value('INDICATOR:RSI').run(timeout=45)
    app.button(key='replay_BTC').click().run(timeout=45)
    assert not app.exception
    assert not app.error
