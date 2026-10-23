"""Verify that published historical measurements trace to executed, frozen artifacts."""
import json,zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
import openpyxl
from indicator_lab.data import ROOT,load
from indicator_lab.runner import fingerprint
from btc_lab.data import sha256
OUT=ROOT/'reports/predictions'
pytestmark=pytest.mark.skipif(not (OUT/'combined.json').exists(),reason='Study not executed yet; never invent saved results')

def study():return json.loads((OUT/'combined.json').read_text())

def test_frozen_code_and_source_provenance_match():
    for m in (5,15,60):
        f=json.loads((OUT/f'frozen_{m}m.json').read_text())
        assert f['code_hashes']==fingerprint()
        assert f['data_manifest_sha256']==sha256((ROOT/'data/predictions/manifest.json').read_bytes())
        assert all(c['passed'] for c in f['causality_checks'])
        assert all(p['model']['training_samples']>=1000 for r in f['models'].values() for p in r['profiles'].values())
    prov=json.loads((ROOT/'sources/tradingview/provenance.json').read_text())
    assert all(sha256((ROOT/k).read_bytes())==v for k,v in prov['files'].items())

def test_all_published_accuracies_recompute_from_exact_saved_forecasts():
    x=study()
    for m in (5,15,60):
        labels=pd.read_csv(OUT/f'labels_{m}m.csv.gz',parse_dates=['forecast_available_at','target_close_at'])
        for r in [v for v in x['rows'] if v['timeframe_minutes']==m]:
            p=np.load(OUT/'probabilities'/f"{r['id']}_{m}m.npz")
            assert len(p['probability'])==len(labels)
            assert p['probability'].dtype==np.float64
            assert np.array_equal(p['predicted_up'],p['probability']>=.5)
            start=pd.Timestamp(r['score_start'],tz='UTC');stop=pd.Timestamp(r['score_stop_exclusive'],tz='UTC')
            mask=(labels.forecast_available_at>=start)&(labels.target_close_at<=stop)&labels.scorable_direction&p['eligible_features']
            assert mask.sum()==r['forecasts']
            measured=float(np.mean(p['predicted_up'][mask]==labels.actual_up[mask].to_numpy()))
            assert measured==r['accuracy']
            assert r['coverage']==r['forecasts']/r['eligible_target_count']
            assert r['coverage']>=.999

def test_rank_only_completed_retrospective_scores_and_full_holm_family():
    x=study();final=[r for r in x['rows'] if r['period']=='retrospective']
    assert x['holm_family_size']==len(final)
    for m in (5,15,60):
        ranked=sorted([r for r in final if r['timeframe_minutes']==m],key=lambda r:r['retrospective_rank'])
        assert [r['retrospective_rank'] for r in ranked]==list(range(1,len(ranked)+1))
        assert all(r['status']=='FORECAST_TEST_COMPLETED' for r in ranked)
        assert all(a['accuracy']>=b['accuracy'] for a,b in zip(ranked,ranked[1:]))
        assert all('NOT AVAILABLE' in r['genuine_untouched_oos_status'] for r in ranked)

def test_increased_costs_cannot_improve_the_same_forecasts():
    x=study();final=[r for r in x['rows'] if r['period']=='retrospective']
    for r in final:
        costs={v['case']:v for v in x['robustness'] if v['id']==r['id'] and v['timeframe_minutes']==r['timeframe_minutes'] and v['kind']=='cost'}
        assert costs['base']['net_return']==r['net_return']
        values=[costs[k]['net_return'] for k in ['half_fee','base','double_fee','double_fee_and_slippage']]
        assert all(a>=b-1e-12 for a,b in zip(values,values[1:]))
        assert len({c['trades'] for c in costs.values()})==1

def test_workbook_is_a_real_excel_file_and_failed_results_are_na():
    p=ROOT/'reports/TradingView_Indicator_Predictions.xlsx';meta=json.loads(p.with_suffix('.json').read_text())
    assert sha256(p.read_bytes())==meta['sha256']
    with zipfile.ZipFile(p) as z:assert z.testzip() is None
    b=openpyxl.load_workbook(p,read_only=True,data_only=True)
    assert {'All Strategies','Top 10','Prediction Accuracy','Charts','Failed Tests','Pine Audit','Walk Forward','Robustness'}<=set(b.sheetnames)
    s=b['All Strategies'];head=[c.value for c in s[4]];data=list(s.iter_rows(min_row=5,values_only=True))
    assert len(data)==441
    status=head.index('status');accuracy=head.index('accuracy')
    for row in data:
        if row[status]!='FORECAST_TEST_COMPLETED':assert row[accuracy]=='N/A'
    assert sum(r[status]=='FORECAST_TEST_COMPLETED' for r in data)==meta['completed_forecast_tests']
    assert b['Pine Audit'].max_row==214
    b.close()

def test_dashboard_switches_between_executed_prediction_timeframes():
    from streamlit.testing.v1 import AppTest
    app=AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
    assert not app.exception
    count=len([r for r in study()['rows'] if r['timeframe_minutes']==5 and r['period']=='retrospective'])
    assert any(len(t.value)==count and 'retrospective_rank' in t.value.columns for t in app.dataframe)
    selector=next(s for s in app.selectbox if s.label=='Prediction chart timeframe')
    selector.select(60).run(timeout=30)
    assert not app.exception
    count=len([r for r in study()['rows'] if r['timeframe_minutes']==60 and r['period']=='retrospective'])
    assert any(len(t.value)==count and 'retrospective_rank' in t.value.columns for t in app.dataframe)
