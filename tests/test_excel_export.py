"""Verify exported facts and Office structure, not just that a file was created."""
from datetime import datetime
import json
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
import pandas as pd
import pytest

from btc_lab.data import DataError, sha256
from btc_lab.strategies import ROOT, load_catalog
from scripts.export_leaderboard import ExcelReport, export, junit_summary, load_inputs


@pytest.fixture(scope='module')
def exported(tmp_path_factory):
    path = tmp_path_factory.mktemp('excel')/'test.xlsx'
    metadata = export(ROOT, path)
    return path, metadata, load_workbook(path, data_only=True)


def table_rows(ws):
    headers = [cell.value for cell in ws[8]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=9, values_only=True) if row[0] is not None]


def test_excel_package_and_required_sheets_are_valid(exported):
    path, metadata, workbook = exported
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        for name in package.namelist():
            if name.endswith('.xml'):
                ET.fromstring(package.read(name))
        assert 'xl/workbook.xml' in package.namelist()
        assert not any('vbaProject' in name for name in package.namelist())
    assert {'All Strategies', 'Top 10', 'Prediction Accuracy', 'Charts', 'Failed Tests'} <= set(workbook.sheetnames)
    assert workbook.active.title == 'Read Me'
    assert metadata['workbook_sha256'] == sha256(path.read_bytes())


def test_all_completed_strategies_and_protocols_are_included_with_correct_metrics(exported):
    _, _, wb = exported
    rows = table_rows(wb['All Strategies'])
    assert len(rows) == 30
    assert len({r['Strategy'] for r in rows}) == 10
    periods = pd.read_csv(ROOT/'reports/audit/period_results.csv')
    legacy = pd.read_csv(ROOT/'reports/summary.csv')
    for r in rows:
        assert r['Execution status'].startswith('COMPLETED')
        if r['Evaluation'].startswith('Exploratory'):
            src = legacy.loc[legacy.strategy == r['Strategy']].iloc[0]
        else:
            src = periods.loc[(periods.strategy == r['Strategy']) & (periods.period == r['Evaluation'])].iloc[0]
        assert r['Net return'] == pytest.approx(src.return_pct/100)
        assert r['Sharpe'] == pytest.approx(src.sharpe_ratio)
        assert r['Max drawdown'] == pytest.approx(src.max_drawdown_pct/100)
        assert r['Win rate'] == pytest.approx(src.win_rate_pct/100)
        assert r['Completed trades'] == src.trades
        assert r['Buy-and-hold return'] == pytest.approx(src.buy_and_hold_return_pct/100)
        assert r['Trading fee per side'] == .001
        assert r['Slippage per side'] == .0005
        assert r['Initial cash (USDT)'] == 10000
        if pd.isna(src.profit_factor):
            assert str(r['Profit factor']).startswith('N/A')
        else:
            assert r['Profit factor'] == pytest.approx(src.profit_factor)


def test_source_links_dates_and_percentage_formats_are_real(exported):
    _, _, wb = exported
    ws = wb['All Strategies']
    headers = {cell.value:cell.column for cell in ws[8]}
    catalog = {s['title']:s for s in load_catalog()}
    for row in table_rows(ws):
        assert row['Original GitHub strategy URL'] == catalog[row['Strategy']]['source_url']
        assert isinstance(row['First input candle UTC'], datetime)
        if row['Evaluation'] == 'nominal_test_2026':
            assert row['First input candle UTC'] == datetime(2026,1,1)
            assert row['Last input candle UTC'] == datetime(2026,9,30,23)
            assert row['Final close deadline UTC'] == datetime(2026,10,1)
            assert row['Last training candle UTC'] == datetime(2024,12,31,22)
            assert row['Warmup candles'] == 300
            assert row['Missing input hours'] == 0
    assert ws.cell(9,headers['Net return']).number_format.startswith('0.00%')
    assert ws.cell(9,headers['Trading fee per side']).number_format == '0.00%'
    assert ws.cell(9,headers['Original GitHub strategy URL']).hyperlink.target.startswith('https://github.com/')
    assert ws.freeze_panes == 'C9'
    assert ws.tables


def test_no_invented_top_ten_or_directional_accuracy(exported):
    _, metadata, wb = exported
    assert metadata['eligible_top_10'] == 0
    assert metadata['forecast_models_scored'] == 0
    for r in table_rows(wb['Top 10']):
        assert r['Genuine unseen rank'] == 'NOT ISSUED'
        assert r['Eligible?'] == 'NO'
    prediction = table_rows(wb['Prediction Accuracy'])
    assert len(prediction) == 10
    for r in prediction:
        assert r['Directional accuracy'].startswith('N/A')
        assert r['Forecasts scored'] == 0
        assert r['Provides forecasts?'] == 'NO'
    for r in table_rows(wb['All Strategies']):
        assert r['Genuinely untouched?'].startswith('NO')
        assert r['Unseen rank'].startswith('NOT ISSUED')
        if r['Evaluation'] != 'Exploratory full period (not OOS)':
            assert r['Passed both 2025 + 2026 gates'] == 'NO'
        if r['Strategy'] == 'Bollinger + RSI' and r['Evaluation']=='validation_2025':
            assert r['OOS validation: 2025 gate'].startswith('PASS')


def test_charts_have_actual_cached_values_and_all_strategies(exported):
    _, _, wb = exported
    charts = wb['Charts']._charts
    assert len(charts) == 4
    assert sorted(len(c.series) for c in charts) == [1,2,4,11]
    chart_data = table_rows(wb['Chart Data'])[:11]
    assert len(chart_data) == 11
    assert chart_data[-1]['Strategy'] == 'Buy and hold (reference)'
    for chart in charts:
        for series in chart.series:
            ref = series.val.numRef
            assert 'Chart Data' in ref.f
            assert ref.numCache.ptCount in (11,273)
            assert len(ref.numCache.pt) == ref.numCache.ptCount
    r = next(r for r in chart_data if r['Strategy']=='Awesome Oscillator + MACD')
    assert r['2026 return'] == pytest.approx(.13589621659774442)


def test_financial_failures_are_not_mislabeled_as_software_errors(exported):
    _, metadata, wb = exported
    rows = table_rows(wb['Failed Tests'])
    assert metadata['software_tests']['failed'] == 0
    assert not any(r['Category']=='Software test' for r in rows)
    assert sum(r['Category']=='Investment validation criterion' for r in rows) == 18
    assert sum(r['Category']=='Combined-period criterion' for r in rows) == 10
    assert any(r['Status']=='MISSING' and '31 missing hours' in r['Observed result'] for r in rows)
    assert any(r['Status']=='UNVERIFIED' for r in rows)


def test_all_source_and_scenario_rows_are_preserved(exported):
    _, _, wb = exported
    assert len(table_rows(wb['Data Sources'])) == 81
    assert len(table_rows(wb['Walk Forward'])) == 30
    assert len(table_rows(wb['Sensitivity'])) == 200
    assert len(table_rows(wb['Training Fits'])) == 120
    for r in table_rows(wb['Data Sources']):
        assert r['Saved verification'] == 'verified'
        assert len(r['Verified archive SHA-256']) == 64
        assert r['Official archive URL'].startswith('https://data.binance.vision/')


def test_unexecuted_rows_are_rejected_before_export(monkeypatch):
    import scripts.export_leaderboard as module
    original = module.read_audit
    def invalid(root):
        frame, meta = original(root)
        frame = frame.copy()
        frame.loc[0,'execution_status'] = 'not_run'
        return frame, meta
    monkeypatch.setattr(module,'read_audit',invalid)
    with pytest.raises(DataError,match='Unexecuted'):
        load_inputs(ROOT)


def test_source_text_cannot_become_an_excel_formula(tmp_path):
    path = tmp_path/'strings.xlsx'
    writer = ExcelReport(path)
    writer.value(writer.sheets['Read Me'],0,0,'=1+1')
    writer.book.close()
    cell = load_workbook(path,data_only=False)['Read Me']['A1']
    assert cell.data_type == 's'
    assert cell.value == '=1+1'


def test_junit_errors_skips_and_passes_are_distinct(tmp_path):
    path = tmp_path/'test.xml'
    path.write_text('<testsuites><testsuite><testcase name="pass"/><testcase name="bad"><failure message="broken"/></testcase><testcase name="skip"><skipped/></testcase></testsuite></testsuites>')
    summary = junit_summary(path)
    assert (summary['tests'],summary['passed'],summary['failed'],summary['skipped']) == (3,1,1,1)
    assert len(summary['failures']) == 1
